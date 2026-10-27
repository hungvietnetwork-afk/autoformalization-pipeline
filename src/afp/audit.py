"""Semantic audit (PDF section 6): 100 compiled samples, two independent raters, Cohen's kappa.

  python -m afp.audit sample --run runs/<run_id> --n 100          # stratified sample + rater sheets
      (the sample is frozen in audit/audit_key.csv once drawn; later calls rebuild the sheets from that key and
       refresh auto_verdict from the run's semantic.jsonl; --redraw draws a new sample)
  python -m afp.audit kappa  --a audit/rater1.xlsx --b audit/rater2.xlsx
  python -m afp.audit kappa  --a audit/llm_rater_A.csv --b audit/llm_rater_B.csv   (pilot with LLM raters)

Labels: DUNG (faithful), GAN_DUNG (minor deviation: a missing side condition, a type that is
slightly off but the mathematical content is kept), SAI (wrong meaning).
Raters only see the natural-language problem, the candidate and (if any) the human reference; the
automatic verdict is stored in audit/audit_key.csv and is NOT shown to them.
"""
from __future__ import annotations

import argparse
import collections
import csv
import json
import os
import random

from .gold import gold_code
from .inference.batchfile import load_units
from .schema import COMPILE_OK

from .paths import AUDIT, CANDIDATES_DB, RAW, ROOT, RUNS, TABLES, TIERS_DIR, UNITS, result_path, run_dir
LABELS = ["DUNG", "GAN_DUNG", "SAI"]
QUOTA = {"proofnet": 20, "minif2f": 20, "math": 45, "gsm8k": 15}

RUBRIC = [
    ("DUNG", "Phát biểu Lean diễn đạt đúng và đủ đề bài: đủ giả thiết, đúng kết luận, đúng kiểu dữ liệu, "
             "với bài có đáp án thì khẳng định đúng đáp án của đúng đại lượng được hỏi."),
    ("GAN_DUNG", "Ý toán học được giữ nhưng có sai lệch nhỏ: thiếu/thừa một điều kiện phụ không ảnh hưởng bản chất, "
                 "kiểu dữ liệu hơi khác (ℕ thay vì ℤ khi mọi giá trị đều dương), cách diễn đạt tương đương nhưng vòng vo."),
    ("SAI", "Sai nghĩa: thiếu giả thiết quan trọng hoặc giả thiết mâu thuẫn, kết luận khác/yếu hơn/tầm thường, "
            "chỉ còn là một phép tính số học thay cho bài toán, sai đáp án hoặc khẳng định sai đại lượng, "
            "sai kiểu làm thay đổi nghĩa (chia/trừ số tự nhiên bị cắt cụt)."),
    ("Chốt sau chạy thử", "Phát biểu chỉ gồm phép tính của lời giải (ví dụ 14 - (2 + 4) = 8), không có biến/giả thiết "
                          "nào đại diện cho các đại lượng của đề, là SAI kể cả khi các con số khớp câu chuyện: "
                          "đúng hay sai lúc này phụ thuộc người formalize đã giải bài, không phải phát biểu."),
]


def sample(run_dir: str, n: int = 100, seed: int = 7) -> list[dict]:
    from .store import CandidateStore, ResultStore
    cfg = json.load(open(os.path.join(run_dir, "config.json")))
    rows = [r for r in ResultStore(os.path.join(run_dir, "run.db")).rows(cfg["run_id"]) if r["status"] == COMPILE_OK]
    cs = CandidateStore(CANDIDATES_DB)
    units = load_units(UNITS)
    sem = {}
    sp = os.path.join(run_dir, "semantic.jsonl")
    if os.path.exists(sp):
        sem = {d["task_hash"]: d for d in map(json.loads, open(sp, encoding="utf-8"))}
    by_src = collections.defaultdict(list)
    for r in rows:
        by_src[units[r["problem_id"]].source].append(r)
    rng = random.Random(seed)
    quota = {s: q for s, q in QUOTA.items()}
    picked = []
    for s, q in quota.items():
        pool = sorted(by_src.get(s, []), key=lambda r: r["task_hash"])
        picked += rng.sample(pool, min(q, len(pool)))
    rest = [r for r in rows if r not in picked]
    if len(picked) < n:      # fill from the remaining pool if a source had too few compiled items
        picked += rng.sample(sorted(rest, key=lambda r: r["task_hash"]), n - len(picked))
    rng.shuffle(picked)
    items = []
    for i, r in enumerate(picked[:n], 1):
        u = units[r["problem_id"]]
        c = cs.get(r["task_hash"])
        items.append({"audit_id": f"A{i:03d}", "problem_id": u.problem_id, "source": u.source, "domain": u.domain,
                      "natural_statement": u.natural_statement, "lean_statement": c.lean_code,
                      "gold_statement": u.metadata.get("gold_formal_statement") or "",
                      "task_hash": r["task_hash"], "auto_verdict": sem.get(r["task_hash"], {}).get("verdict", "")})
    return items


def items_from_key(run_dir: str, key_path: str = None) -> list[dict]:
    """Rebuild the audit items of a frozen sample (audit_key.csv), refreshing auto_verdict from the run."""
    from .store import CandidateStore
    key_path = key_path or os.path.join(AUDIT, "audit_key.csv")
    cs = CandidateStore(CANDIDATES_DB)
    units = load_units(UNITS)
    sem = {}
    sp = os.path.join(run_dir, "semantic.jsonl")
    if os.path.exists(sp):
        sem = {d["task_hash"]: d for d in map(json.loads, open(sp, encoding="utf-8"))}
    items = []
    for r in csv.DictReader(open(key_path, encoding="utf-8")):
        u, c = units[r["problem_id"]], cs.get(r["task_hash"])
        items.append({"audit_id": r["audit_id"], "problem_id": u.problem_id, "source": u.source, "domain": u.domain,
                      "natural_statement": u.natural_statement, "lean_statement": c.lean_code,
                      "gold_statement": u.metadata.get("gold_formal_statement") or "", "task_hash": r["task_hash"],
                      "auto_verdict": sem.get(r["task_hash"], {}).get("verdict", "") or r.get("auto_verdict", "")})
    return items


def export(items: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    os.makedirs(AUDIT, exist_ok=True)
    with open(os.path.join(AUDIT, "audit_key.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["audit_id", "problem_id", "source", "domain", "task_hash", "auto_verdict"])
        w.writeheader()
        for it in items:
            w.writerow({k: it[k] for k in w.fieldnames})
    with open(os.path.join(AUDIT, "audit_items.txt"), "w", encoding="utf-8") as f:   # plain text for LLM raters
        for it in items:
            f.write(f"### {it['audit_id']}\n[PROBLEM]\n{it['natural_statement']}\n[LEAN]\n{it['lean_statement']}\n")
            if it["gold_statement"]:
                f.write(f"[REFERENCE (human, older Mathlib)]\n{it['gold_statement']}\n")
            f.write("\n")
    for rater in ("rater1", "rater2"):
        path = os.path.join(AUDIT, f"{rater}.xlsx")
        if os.path.exists(path) and _read_labels(path):
            print(f"{path} already holds labels - not overwritten")
            continue
        wb = Workbook()
        ws = wb.active
        ws.title = "audit"
        head = ["audit_id", "source", "de_bai (NL)", "phat_bieu_Lean (ung vien)", "tham_khao_gold", "nhan", "ghi_chu"]
        ws.append(head)
        for it in items:
            ws.append([it["audit_id"], it["source"], it["natural_statement"], it["lean_statement"],
                       it["gold_statement"], "", ""])
        for col, wdt in zip("ABCDEFG", [9, 10, 60, 70, 50, 12, 30]):
            ws.column_dimensions[col].width = wdt
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"F2:F{len(items) + 1}")
        ws.freeze_panes = "C2"
        rs = wb.create_sheet("huong_dan")
        rs.append(["Nhãn", "Ý nghĩa"])
        for lab, desc in RUBRIC:
            rs.append([lab, desc])
        rs.append([])
        rs.append(["Quy tắc", "Chấm độc lập, không trao đổi với người chấm còn lại. Không xem kết quả kiểm tra tự động. "
                              "Cột tham_khao_gold là phát biểu do người viết cho Mathlib cũ, chỉ để tham khảo."])
        rs.column_dimensions["A"].width, rs.column_dimensions["B"].width = 12, 120
        for row in rs.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(path)
    print(f"exported {len(items)} items to {AUDIT}")


def _read_labels(path: str) -> dict[str, str]:
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path)["audit"]
        out = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[5]:
                out[str(row[0])] = str(row[5]).strip().upper()
        return out
    with open(path, encoding="utf-8") as f:
        return {r["audit_id"]: r["label"].strip().upper() for r in csv.DictReader(f) if r.get("label")}


def kappa(a_path: str, b_path: str) -> dict:
    from sklearn.metrics import cohen_kappa_score, confusion_matrix
    a, b = _read_labels(a_path), _read_labels(b_path)
    ids = sorted(set(a) & set(b))
    la, lb = [a[i] for i in ids], [b[i] for i in ids]
    bin_a = ["OK" if x == "DUNG" else "NOT_OK" for x in la]
    bin_b = ["OK" if x == "DUNG" else "NOT_OK" for x in lb]
    bin2_a = ["WRONG" if x == "SAI" else "ACCEPTABLE" for x in la]
    bin2_b = ["WRONG" if x == "SAI" else "ACCEPTABLE" for x in lb]
    rep = {
        "n_items": len(ids),
        "observed_agreement_3class": round(sum(x == y for x, y in zip(la, lb)) / len(ids), 4),
        "kappa_3class": round(cohen_kappa_score(la, lb, labels=LABELS), 4),
        "kappa_3class_linear_weighted": round(cohen_kappa_score(la, lb, labels=LABELS, weights="linear"), 4),
        "kappa_binary_DUNG_vs_rest": round(cohen_kappa_score(bin_a, bin_b), 4),
        "kappa_binary_SAI_vs_rest": round(cohen_kappa_score(bin2_a, bin2_b), 4),
        "label_dist_a": dict(collections.Counter(la)), "label_dist_b": dict(collections.Counter(lb)),
        "confusion_3class(rows=a, cols=b)": confusion_matrix(la, lb, labels=LABELS).tolist(),
    }
    return rep


AUTO_ERROR = {"VACUOUS", "FALSE", "TRIVIAL", "ANSWER_MISSING", "ARITHMETIC_ONLY"}


def compare_auto(rater_path: str, key_path: str = None, sem_paths: list[str] | None = None) -> dict:
    """How well do the automatic semantic flags agree with a rater? Reported twice: against the rater's SAI labels
    and against everything the rater did not label DUNG (SAI or GAN_DUNG)."""
    key_path = key_path or os.path.join(AUDIT, "audit_key.csv")
    labels = _read_labels(rater_path)
    key = {r["audit_id"]: r for r in csv.DictReader(open(key_path, encoding="utf-8"))}
    verdicts = {}
    for sp in sem_paths or []:
        for line in open(sp, encoding="utf-8"):
            d = json.loads(line)
            verdicts[d["task_hash"]] = d["verdict"]

    def score(is_wrong) -> dict:
        tp = fp = fn = tn = 0
        for aid, lab in labels.items():
            v = verdicts.get(key[aid]["task_hash"]) or key[aid].get("auto_verdict", "")
            flag, wrong = v in AUTO_ERROR, is_wrong(lab)
            tp += flag and wrong
            fp += flag and not wrong
            fn += (not flag) and wrong
            tn += (not flag) and (not wrong)
        return {"n": tp + fp + fn + tn, "auto_flagged": tp + fp, "rater_positive": tp + fn, "tp": tp, "fp": fp,
                "fn": fn, "tn": tn, "precision": round(tp / (tp + fp), 3) if tp + fp else None,
                "recall": round(tp / (tp + fn), 3) if tp + fn else None}

    flagged = collections.Counter()
    for aid in labels:
        flagged[verdicts.get(key[aid]["task_hash"]) or key[aid].get("auto_verdict", "") or "none"] += 1
    return {"rater": os.path.basename(rater_path), "auto_error_verdicts": sorted(AUTO_ERROR),
            "verdicts_on_sample": dict(flagged), "vs_SAI": score(lambda lab: lab == "SAI"),
            "vs_not_DUNG": score(lambda lab: lab != "DUNG")}


def main(argv=None):
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample")
    s.add_argument("--run", required=True)
    s.add_argument("--n", type=int, default=100)
    s.add_argument("--redraw", action="store_true", help="draw a new sample even if audit/audit_key.csv exists")
    k = sub.add_parser("kappa")
    k.add_argument("--a", required=True)
    k.add_argument("--b", required=True)
    k.add_argument("--out", help="also write the report to this JSON file")
    c = sub.add_parser("compare", help="automatic semantic flags vs a rater's SAI labels")
    c.add_argument("--rater", required=True)
    c.add_argument("--sem", nargs="+", required=True, help="semantic.jsonl of the runs the items come from")
    c.add_argument("--out")
    a = ap.parse_args(argv)
    if a.cmd == "sample":
        frozen = os.path.exists(os.path.join(AUDIT, "audit_key.csv")) and not a.redraw
        export(items_from_key(run_dir(a.run)) if frozen else sample(run_dir(a.run), a.n))
        return
    rep = kappa(a.a, a.b) if a.cmd == "kappa" else compare_auto(a.rater, sem_paths=a.sem)
    if a.out:
        json.dump(rep, open(a.out, "w"), indent=1, ensure_ascii=False)
    print(json.dumps(rep, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
